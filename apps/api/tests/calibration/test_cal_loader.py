import openpyxl
import pandas as pd
import pytest

from ite_api.calibration.cal_loader import load_calibration_sheet, load_workbook


def test_load_workbook_returns_sheet_names(workbook_xlsx):
    wb, names = load_workbook(workbook_xlsx)
    assert len(names) > 0


def test_load_calibration_sheet_returns_typed_dataframe(workbook_xlsx):
    wb, names = load_workbook(workbook_xlsx)
    df = load_calibration_sheet(wb, names[0])
    assert isinstance(df, pd.DataFrame)
    assert set(df.columns) == {"timestamp", "temp"}
    assert len(df) > 0
    assert pd.api.types.is_datetime64_any_dtype(df["timestamp"])
    assert pd.api.types.is_float_dtype(df["temp"])


def test_load_calibration_sheet_raises_when_every_row_fails_to_parse():
    """A sheet whose data rows have no usable timestamp/temp (e.g. wrong columns,
    a re-shuffled export) must raise, not silently return an empty DataFrame —
    an empty-but-successful load flowed straight into the matcher as "no data",
    masking a real malformed-workbook problem."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LOGGER001"
    ws.append(["idx", "timestamp", "unused", "temp"])
    ws.append([1, None, "x", None])
    ws.append([2, None, "x", None])
    with pytest.raises(ValueError, match="no valid"):
        load_calibration_sheet(wb, "LOGGER001")
