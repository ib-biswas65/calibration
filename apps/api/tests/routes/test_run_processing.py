"""Integration tests for _do_process's per-logger match-failure handling.

Exercises the pipeline directly against a real Postgres (via db_session), bypassing
the app/TestClient — no HTTP layer needed for this behavior.
"""

from pathlib import Path

import openpyxl
from sqlalchemy import select

from ite_api.config import Settings
from ite_api.db.models.calibration import CalibrationRun, LoggerResult
from ite_api.routes.runs import _do_process, _run_detail


def _make_run(db_session, **overrides):
    defaults = dict(
        batch_name="Partial Failure Test",
        status="processing",
        testing_start="2026-04-14T09:00:00+00:00",
        testing_end="2026-04-14T10:00:00+00:00",
        certificate_date="2026-04-15",
        threshold_c=0.5,
        setpoints=[
            {
                "target_c": 5.0,
                "start_at": "2026-04-14T09:00:00+00:00",
                "end_at": "2026-04-14T10:00:00+00:00",
            }
        ],
        start_cert_no="0000009000",
        cert_width=10,
        test_date_jp="2026年4月14日",
        doc_date_jp="2026年4月15日",
    )
    defaults.update(overrides)
    run = CalibrationRun(**defaults)
    db_session.add(run)
    db_session.commit()
    db_session.refresh(run)
    return run


def _write_reference_csv(tmp_path: Path) -> Path:
    p = tmp_path / "ref.csv"
    p.write_text(
        "2026/04/14 09:01:00,5.0\n"
        "2026/04/14 09:02:00,5.05\n"
    )
    return p


def _write_workbook(tmp_path: Path) -> Path:
    """One sheet with a calibration reading close to the reference (should match),
    one sheet with a reading nowhere near any reference value in the window
    (must NOT fabricate a match — this is the regression case for the matcher bug)."""
    wb = openpyxl.Workbook()
    good = wb.active
    good.title = "GOOD001"
    good.append(["idx", "timestamp", "unused", "temp"])
    good.append([1, "2026-04-14 09:01:30", "x", 5.1])

    bad = wb.create_sheet("BAD001")
    bad.append(["idx", "timestamp", "unused", "temp"])
    bad.append([1, "2026-04-14 09:01:30", "x", 99.0])

    p = tmp_path / "wb.xlsx"
    wb.save(p)
    return p


def test_do_process_marks_unmatched_logger_invalid_and_run_partial(db_session, tmp_path):
    run = _make_run(db_session)
    ref_path = _write_reference_csv(tmp_path)
    wb_path = _write_workbook(tmp_path)
    settings = Settings(data_dir=tmp_path / "data")

    _do_process(run, [ref_path], wb_path, settings, db_session)

    db_session.refresh(run)
    assert run.status == "partial"

    results = {
        r.sheet_name: r
        for r in db_session.scalars(select(LoggerResult).where(LoggerResult.run_id == run.id)).all()
    }

    good = results["GOOD001"]
    assert good.verdict in ("pass", "fail")
    assert good.cert_path is not None
    assert good.failure_reason is None

    bad = results["BAD001"]
    assert bad.verdict == "invalid"
    assert bad.cert_path is None
    assert bad.failure_reason is not None
    assert bad.max_deviation_c is None

    detail = _run_detail(run, db_session)
    bad_serialized = next(r for r in detail.results if r["sheet_name"] == "BAD001")
    assert bad_serialized["failure_reason"] == bad.failure_reason
    good_serialized = next(r for r in detail.results if r["sheet_name"] == "GOOD001")
    assert good_serialized["failure_reason"] is None


def test_do_process_all_loggers_invalid_marks_run_failed(db_session, tmp_path):
    """If every logger in the run fails to match, the run is failed outright —
    "partial" would be misleading when nothing actually succeeded."""
    wb = openpyxl.Workbook()
    only = wb.active
    only.title = "ONLYBAD"
    only.append(["idx", "timestamp", "unused", "temp"])
    only.append([1, "2026-04-14 09:01:30", "x", 99.0])
    wb_path = tmp_path / "wb.xlsx"
    wb.save(wb_path)

    run = _make_run(db_session)
    ref_path = _write_reference_csv(tmp_path)
    settings = Settings(data_dir=tmp_path / "data")

    _do_process(run, [ref_path], wb_path, settings, db_session)

    db_session.refresh(run)
    assert run.status == "failed"
    results = db_session.scalars(select(LoggerResult).where(LoggerResult.run_id == run.id)).all()
    assert len(results) == 1
    assert results[0].verdict == "invalid"
