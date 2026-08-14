"""Calibration orchestrator: load → match → fill → save."""

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document

from ite_api.calibration.cal_loader import load_calibration_sheet, load_workbook
from ite_api.calibration.docx_filler import (
    fill_results_table,
    replace_text_everywhere,
)
from ite_api.calibration.matcher import find_values_for_target
from ite_api.calibration.ref_loader import combine_refs, load_ref_auto


@dataclass(frozen=True)
class SetpointWindow:
    target: float
    start: datetime
    end: datetime


@dataclass(frozen=True)
class RunConfig:
    cert_no: str           # new cert number to write
    serial: str            # new logger serial to write
    test_date_jp: str      # e.g. "2026年4月14日"
    doc_date_jp: str       # e.g. "2026年4月15日"
    template_path: Path
    output_dir: Path
    setpoints: list[SetpointWindow]
    # Placeholder strings present in the shipped template.docx
    template_cert_no: str = "0000001700"
    template_serial: str = "190124110002417"
    template_test_date: str = "2026年3月4日"
    template_doc_date: str = "2026年3月6日"


def run_one_logger(cfg: RunConfig, *, sheet_name: str, wb, ref_df) -> Path:
    """Generate one certificate for one sheet. Returns the written .docx path."""
    cal_df = load_calibration_sheet(wb, sheet_name)
    ordered_values: list[tuple[float | None, float | None]] = []
    for sp in cfg.setpoints:
        ref_v, cal_v, _ = find_values_for_target(
            cal_df, ref_df, sp.target, sp.start, sp.end
        )
        ordered_values.append((ref_v, cal_v))

    doc = Document(str(cfg.template_path))
    replace_text_everywhere(doc, {
        cfg.template_cert_no: cfg.cert_no,
        cfg.template_serial: cfg.serial,
        cfg.template_test_date: cfg.test_date_jp,
        cfg.template_doc_date: cfg.doc_date_jp,
    })
    fill_results_table(doc, ordered_values)

    cfg.output_dir.mkdir(parents=True, exist_ok=True)
    out = cfg.output_dir / f"Calibration_Certificate_{cfg.cert_no}_{cfg.serial}.docx"
    doc.save(str(out))
    return out


@dataclass(frozen=True)
class BatchConfig:
    start_cert_no: str
    cert_width: int
    test_date_jp: str
    doc_date_jp: str
    template_path: Path
    calibration_xlsxs: list[Path]   # one or more workbooks; sheets concatenated in order
    reference_csvs: list[Path]      # one or more reference CSVs; rows concatenated
    output_dir: Path
    setpoints: list[SetpointWindow]
    serial_from_sheet: bool = True


def _format_cert_no(n: int, width: int) -> str:
    return str(n).zfill(width)


def run_calibration(cfg: BatchConfig) -> list[Path]:
    """Process every sheet across all workbooks. Returns written paths in iteration order."""
    ref_df = combine_refs([load_ref_auto(p) for p in cfg.reference_csvs])
    start = int(cfg.start_cert_no)
    written: list[Path] = []
    cert_idx = 0
    for xlsx_path in cfg.calibration_xlsxs:
        wb, sheet_names = load_workbook(xlsx_path)
        for name in sheet_names:
            cert_no = _format_cert_no(start + cert_idx, cfg.cert_width)
            serial = name.strip() if cfg.serial_from_sheet else ""
            run_cfg = RunConfig(
                cert_no=cert_no,
                serial=serial,
                test_date_jp=cfg.test_date_jp,
                doc_date_jp=cfg.doc_date_jp,
                template_path=cfg.template_path,
                output_dir=cfg.output_dir,
                setpoints=cfg.setpoints,
            )
            written.append(run_one_logger(run_cfg, sheet_name=name, wb=wb, ref_df=ref_df))
            cert_idx += 1
    return written
